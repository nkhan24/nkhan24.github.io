"""Make the aero-map charts for the Williams sim-racing page.

Reads the cached values from the F1 25 aero-map workbook (keep the workbook
OUTSIDE the repo) and writes two WebP charts.

Run from the repo root:
    .venv/Scripts/python.exe tools/make_aero_charts.py ../portfolio-drafts/data/F1AeromapsTesting_F1_25_Check.xlsx assets/images/projects
"""
import io
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from PIL import Image

# Site colours (css/style.css tokens) and a dark categorical palette
# checked for colour-blind separation against the card surface.
SURFACE, TEXT, MUTED, GRID = "#171a21", "#e8eaed", "#9aa0a6", "#2a2f3a"
BLUE, ORANGE, AQUA, YELLOW = "#3987e5", "#d95926", "#199e70", "#c98500"


def read_rows(wb, sheet, rows):
    """Return x (setting), front load, rear load, total load, aero balance."""
    ws = wb[sheet]
    out = []
    for r in rows:
        x = ws[f"F{r}"].value
        out.append({
            "x": 0 if isinstance(x, str) or x is None else x,
            "front": ws[f"O{r}"].value,
            "rear": ws[f"P{r}"].value,
            "total": ws[f"Q{r}"].value,
            "ab": ws[f"X{r}"].value,
        })
    return out


def style_axes(ax):
    ax.set_facecolor(SURFACE)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(GRID)
    ax.tick_params(colors=MUTED, labelsize=13)
    ax.grid(color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)


def save_webp(fig, dest):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=160, facecolor=SURFACE)
    plt.close(fig)
    buf.seek(0)
    Image.open(buf).convert("RGB").save(dest, "WEBP", lossless=True, method=6)
    print(f"saved  {dest}")


src, out = Path(sys.argv[1]), Path(sys.argv[2])
wb = openpyxl.load_workbook(src, data_only=True)

mex_wing = "Sweeps Wing Changes MEX"
mex_rh = "Sweeps RH Changes MEX"
bhn_wing = "Sweeps Wing Changes (BHN)"

fw = read_rows(wb, mex_wing, [7] + list(range(11, 21)))     # front wing 0 to 50
rw = read_rows(wb, mex_wing, [24] + list(range(28, 38)))    # rear wing 0 to 50
frh = read_rows(wb, mex_rh, [17] + list(range(18, 25)))     # front ride height 15 to 35
rrh = read_rows(wb, mex_rh, [6] + list(range(7, 14)))       # rear ride height 40 to 60
fw_b = read_rows(wb, bhn_wing, [7] + list(range(11, 21)))
rw_b = read_rows(wb, bhn_wing, [24] + list(range(28, 38)))  # row 38 left out

plt.rcParams["font.family"] = "sans-serif"

# Chart 1: which way each setup change moves the car
base = fw[0]
fig, ax = plt.subplots(figsize=(10, 5.625))
fig.patch.set_facecolor(SURFACE)
style_axes(ax)
levers = [
    (fw, BLUE, "o", "Front wing +50"),
    (rw, ORANGE, "s", "Rear wing +50"),
    (frh, AQUA, "^", "Front ride height +20"),
    (rrh, YELLOW, "D", "Rear ride height +20"),
]
for data, colour, marker, label in levers:
    xs = [d["ab"] for d in data]
    ys = [d["total"] - base["total"] for d in data]
    ax.plot(xs, ys, color=colour, linewidth=2, marker=marker, markersize=5,
            markeredgecolor=SURFACE, markeredgewidth=1, label=label)
    ax.annotate("", xy=(xs[-1], ys[-1]), xytext=(xs[-2], ys[-2]),
                arrowprops=dict(arrowstyle="-|>", color=colour, lw=2))
    ax.annotate(label, (xs[-1], ys[-1]), textcoords="offset points",
                xytext=(8, -4 if ys[-1] < 0 else 6), color=TEXT, fontsize=13)
ax.plot(base["ab"], 0, marker="o", markersize=9, color=TEXT, zorder=5)
ax.annotate("Baseline", (base["ab"], 0), textcoords="offset points",
            xytext=(-80, 12), color=TEXT, fontsize=13)
ax.axhline(0, color=MUTED, linewidth=0.8, linestyle=":")
ax.set_xlabel("Aero balance (% of downforce on the front axle)", color=MUTED, fontsize=14)
ax.set_ylabel("Change in total downforce at 200 km/h (N)", color=MUTED, fontsize=14)
ax.set_xlim(33.5, 51.5)
ax.legend(loc="lower right", frameon=False, labelcolor=TEXT, fontsize=12)
fig.tight_layout()
save_webp(fig, out / "wfg-aero-direction.webp")

# Chart 2: repeatability, Mexico against a Bahrain repeat
fig, ax = plt.subplots(figsize=(10, 5.625))
fig.patch.set_facecolor(SURFACE)
style_axes(ax)
for data, colour, marker, name, track, dashed in [
    (fw, BLUE, "o", "Front wing", "Mexico", False),
    (fw_b, BLUE, "o", "Front wing", "Bahrain", True),
    (rw, ORANGE, "s", "Rear wing", "Mexico", False),
    (rw_b, ORANGE, "s", "Rear wing", "Bahrain", True),
]:
    ax.plot([d["x"] for d in data], [d["ab"] for d in data], color=colour,
            linewidth=2, linestyle="--" if dashed else "-", marker=marker,
            markersize=6, markerfacecolor=SURFACE if dashed else colour,
            markeredgecolor=colour, label=f"{name}, {track}")
ax.annotate("Front wing", (50, fw[-1]["ab"]), textcoords="offset points",
            xytext=(8, -4), color=TEXT, fontsize=13)
ax.annotate("Rear wing", (50, rw[-1]["ab"]), textcoords="offset points",
            xytext=(8, -4), color=TEXT, fontsize=13)
ax.set_xlabel("Wing setting (clicks)", color=MUTED, fontsize=14)
ax.set_ylabel("Aero balance (% front)", color=MUTED, fontsize=14)
ax.set_xlim(-2, 62)
ax.legend(loc="upper left", frameon=False, labelcolor=TEXT, fontsize=12, ncol=2)
fig.tight_layout()
save_webp(fig, out / "wfg-aero-repeat.webp")